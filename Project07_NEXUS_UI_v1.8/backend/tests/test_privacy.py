import base64
import json
import io
import csv
import os
from pathlib import Path
import pandas as pd
import pytest
from cryptography.exceptions import InvalidTag
from services import privacy


@pytest.fixture
def vault(tmp_path, monkeypatch):
    monkeypatch.setenv('PROJECT07_VAULT_MASTER_KEY', base64.b64encode(os.urandom(32)).decode())
    monkeypatch.setenv('PROJECT07_VAULT_DIR', str(tmp_path / 'vault'))
    return tmp_path / 'vault'


@pytest.mark.parametrize('filename,raw', [('customer.csv',b'full_name,phone\nSecret Person,0901234567'),
    ('document.pdf',b'%PDF binary archive'), ('unknown.bin',bytes(range(256)))])
def test_encrypt_restore_no_plaintext(vault, filename, raw):
    protected = privacy.protect_file(raw, filename)
    assert filename.encode() not in protected['encrypted_bytes']
    assert raw not in protected['encrypted_bytes']
    assert privacy.restore_file(protected['encrypted_bytes']) == (raw, filename)
    for path in vault.rglob('*'):
        if path.is_file():
            assert raw not in path.read_bytes()
            if os.name == 'posix': assert path.stat().st_mode & 0o777 == 0o600
        elif os.name == 'posix': assert path.stat().st_mode & 0o777 == 0o700


def test_tamper_detected(vault):
    result = privacy.protect_file(b'secret', 'source.csv')
    obj = json.loads(result['encrypted_bytes'])
    ciphertext = bytearray(privacy.unb64(obj['ciphertext'])); ciphertext[0] ^= 1
    obj['ciphertext'] = privacy.b64(ciphertext)
    with pytest.raises(InvalidTag): privacy.restore_file(json.dumps(obj).encode())


def test_wrong_master_detected(vault, monkeypatch):
    result = privacy.protect_file(b'secret', 'source.csv')
    monkeypatch.setenv('PROJECT07_VAULT_MASTER_KEY', base64.b64encode(os.urandom(32)).decode())
    with pytest.raises(InvalidTag): privacy.restore_file(result['encrypted_bytes'])


def test_missing_master_does_not_write(tmp_path, monkeypatch):
    monkeypatch.delenv('PROJECT07_VAULT_MASTER_KEY', raising=False)
    monkeypatch.setenv('PROJECT07_VAULT_DIR', str(tmp_path / 'never-created'))
    with pytest.raises(privacy.PrivacyError): privacy.protect_file(b'data', 'source.csv')
    assert not (tmp_path / 'never-created').exists()


def test_unknown_columns_removed_and_personal_columns_tokenized(vault):
    frame = pd.DataFrame({'full_name':['Secret Person'], 'phone':['0901234567'],
                          'amount_vnd':['100000'], 'note':['private unstructured data']})
    plan = privacy.suggest_columns(frame)
    assert {x['column']: x['action'] for x in plan} == {'full_name':'TOKENIZE','phone':'TOKENIZE','amount_vnd':'KEEP','note':'MASK'}
    encrypted = privacy.protect_file(frame.to_csv(index=False).encode(), 'source.csv')
    result = privacy.pseudonymize(frame, plan, 'SCOPE', 'Test reviewer', encrypted['file_id'])
    assert b'Secret Person' not in result['csv_bytes'] and b'0901234567' not in result['csv_bytes']
    assert result['frame']['note'][0] == '[REDACTED]' and result['frame']['amount_vnd'][0] == '100000'
    assert b'Secret Person' in result['audit_mapping_csv_bytes']
    assert not list((vault / 'maps').iterdir())
    assert list(result['frame'].columns) == list(frame.columns)


def test_scope_keeps_paired_employee_ids_consistent(vault):
    master = privacy.master_key()
    assert privacy.token_value(master, 'audit-1', 'EMP01') == privacy.token_value(master, 'audit-1', 'EMP01')
    assert privacy.token_value(master, 'audit-1', 'EMP01') != privacy.token_value(master, 'audit-2', 'EMP01')
    frames = [pd.DataFrame({'maker_id':['EMP01'], 'approver_id':['EMP02']}), pd.DataFrame({'employee_id':['EMP01']})]
    protected = [privacy.protect_file(f.to_csv(index=False).encode(), 'source.csv') for f in frames]
    outputs = [privacy.pseudonymize(f, privacy.suggest_columns(f), 'scope', 'reviewer', p['file_id'])['frame'] for f,p in zip(frames, protected)]
    assert outputs[0]['maker_id'][0] == outputs[1]['employee_id'][0]
    assert outputs[0]['maker_id'][0] != outputs[0]['approver_id'][0]


def test_no_unreviewed_pseudonym_export(vault):
    frame = pd.DataFrame({'customer_id':['customer']})
    protected = privacy.protect_file(b'data', 'source')
    with pytest.raises(privacy.PrivacyError):
        privacy.pseudonymize(frame, privacy.suggest_columns(frame), 'scope', '', protected['file_id'])
    assert not list((vault / 'maps').iterdir())


def test_formula_escape_and_filename_safety(vault):
    frame = pd.DataFrame({'name':['=HYPERLINK("bad")']})
    protected = privacy.protect_file(b'source', '../../secret.csv')
    assert privacy.restore_file(protected['encrypted_bytes'])[1] == 'secret.csv'
    result = privacy.pseudonymize(frame, [{'column':'name','action':'KEEP'}], 'scope', 'reviewer', protected['file_id'])
    assert "'=HYPERLINK" in result['csv_bytes'].decode('utf-8-sig')


def test_invalid_actions_and_path_ids_rejected(vault):
    frame = pd.DataFrame({'customer_id':['C01']})
    protected = privacy.protect_file(b'data', 'source')
    with pytest.raises(privacy.PrivacyError):
        privacy.pseudonymize(frame, [{'column':'customer_id','action':'BOGUS'}], 'scope', 'reviewer', protected['file_id'])
    with pytest.raises(privacy.PrivacyError): privacy.valid_id('../outside')


def test_invalid_or_unsupported_tables_rejected(vault):
    with pytest.raises(privacy.PrivacyError): privacy.read_table(b'data', 'source.pdf')
    with pytest.raises(privacy.PrivacyError): privacy.read_table(b'not an excel file', 'source.xlsx')


def test_missing_wrapped_key_prevents_recovery(vault):
    protected = privacy.protect_file(b'secret', 'source')
    (vault / 'keys' / (protected['file_id'] + '.key.enc')).unlink()
    with pytest.raises(FileNotFoundError): privacy.restore_file(protected['encrypted_bytes'])


def test_csv_preserves_columns_delimiter_bom_and_newlines(vault):
    raw = '\ufeffrow_id;full_name;amount_vnd;note\r\nR1;Nguyễn An;100000;private text\r\n'.encode('utf-8')
    frame = privacy.read_table(raw, 'input.csv')
    protected = privacy.protect_file(raw, 'input.csv')
    result = privacy.pseudonymize(frame, privacy.suggest_columns(frame), 'scope', 'auditor', protected['file_id'], raw, 'input.csv')
    assert result['export_filename'] == 'protected_input.csv'
    assert result['export_bytes'].startswith(b'\xef\xbb\xbf')
    text = result['export_bytes'].decode('utf-8-sig')
    assert text.split('\r\n')[0] == 'row_id;full_name;amount_vnd;note'
    assert '\n' not in text.replace('\r\n','')
    data = list(csv.reader(io.StringIO(text), delimiter=';'))
    assert len(data[1]) == 4 and data[1][2] == '100000' and data[1][3] == '[REDACTED]'
    assert not list((vault / 'maps').iterdir())
    report = json.loads(result['audit_report_json_bytes'])
    assert report['column_count'] == 4 and len(report['cell_details']) == 4
    assert report['cell_details'][1]['original_value'] == 'Nguyễn An'


def test_legacy_remove_masks_in_place(vault):
    frame = pd.DataFrame({'customer_id':['C01'], 'private_notes':['secret']})
    file = privacy.protect_file(b'source', 'input.csv')
    result = privacy.pseudonymize(frame,[{'column':'customer_id','action':'TOKENIZE'},
                                      {'column':'private_notes','action':'REMOVE'}],'scope','auditor',file['file_id'])
    assert list(result['frame'].columns) == list(frame.columns)
    assert result['frame']['private_notes'][0] == '[REDACTED]'


def test_xlsx_preserves_single_sheet_layout_and_numeric_types(vault):
    # Synthetic in-memory fixture for application serializer behavior, not a user deliverable.
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Loans'
    sheet.append(['full_name','amount_vnd','private_note']); sheet.append(['Secret Person',125000000,'secret'])
    sheet['A2'].font = Font(bold=True, color='FF123456')
    sheet['B2'].number_format = '#,##0'; sheet.column_dimensions['A'].width = 32
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = 'A1:C2'
    buffer = io.BytesIO(); workbook.save(buffer); raw = buffer.getvalue()
    frame = privacy.read_table(raw,'input.xlsx'); file = privacy.protect_file(raw,'input.xlsx')
    result = privacy.pseudonymize(frame,privacy.suggest_columns(frame),'scope','auditor',file['file_id'],raw,'input.xlsx')
    output = load_workbook(io.BytesIO(result['export_bytes'])); sheet = output['Loans']
    assert output.sheetnames == ['Loans'] and sheet.max_column == 3 and sheet.max_row == 2
    assert sheet['A2'].value.startswith('NX_') and sheet['A2'].font.bold and sheet['A2'].font.color.rgb == 'FF123456'
    assert sheet['B2'].value == 125000000 and sheet['B2'].data_type == 'n' and sheet['B2'].number_format == '#,##0'
    assert sheet['C2'].value == '[REDACTED]'
    assert sheet.column_dimensions['A'].width == 32 and sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == 'A1:C2'
    assert result['export_filename'].endswith('.xlsx')
    assert not list((vault / 'maps').iterdir())


def test_multi_sheet_export_fails_before_releasing_mapping(vault):
    from openpyxl import Workbook
    workbook=Workbook(); workbook.active.append(['full_name']); workbook.active.append(['Person'])
    workbook.create_sheet('Hidden personal data')
    buffer=io.BytesIO();workbook.save(buffer);raw=buffer.getvalue()
    frame=privacy.read_table(raw,'multi.xlsx');file=privacy.protect_file(raw,'multi.xlsx')
    with pytest.raises(privacy.PrivacyError,match='nhiều sheet'):
        privacy.pseudonymize(frame,privacy.suggest_columns(frame),'scope','auditor',file['file_id'],raw,'multi.xlsx')
    assert not list((vault / 'maps').iterdir())
