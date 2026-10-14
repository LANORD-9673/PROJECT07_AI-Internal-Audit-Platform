import io
from docx import Document
from openpyxl import Workbook
import fitz
from test_imports import LOANS, AUTH


def _policy(client, filename, content):
    return client.post('/api/import/policies', data={'policy_id':'NEXUS-TEST','policy_set_id':'nexus_v1_1','effective_from':'2026-10-08T00:00:00+07:00'}, files={'file':(filename, content)})


def test_docx_and_extraction(client):
    doc=Document(); doc.add_paragraph('Quy trinh phe duyet khoan vay')
    table=doc.add_table(rows=1, cols=2); table.cell(0,0).text='Ma'; table.cell(0,1).text='Gia tri'
    stream=io.BytesIO(); doc.save(stream)
    r=_policy(client,'nexus.docx',stream.getvalue()); assert r.status_code==201, r.text
    sid=r.json()['data']['snapshot_id']
    t=client.get('/api/policies/'+sid+'/text'); assert t.status_code==200
    assert 'Quy trinh phe duyet' in t.json()['data']['extracted_text']
    assert 'Ma | Gia tri' in t.json()['data']['extracted_text']


def test_pdf_and_scan_rejection(client):
    doc=fitz.open(); page=doc.new_page(); page.insert_text((70,70),'Policy approval authority')
    r=_policy(client,'p.pdf',doc.tobytes()); assert r.status_code==201, r.text
    doc2=fitz.open(); doc2.new_page()
    bad=_policy(client,'scan.pdf',doc2.tobytes()); assert bad.status_code==422
    assert bad.json()['errors'][0]['code']=='NO_EXTRACTABLE_TEXT'


def test_xlsx_loans_and_authority(client):
    for path, field, name, csv in [('/api/import/loans','dataset_id','loans.xlsx',LOANS),('/api/import/authorities','authority_id','auth.xlsx',AUTH)]:
        wb=Workbook(); ws=wb.active
        for line in csv.strip().splitlines(): ws.append(line.split(','))
        stream=io.BytesIO(); wb.save(stream)
        r=client.post(path,data={field:'xlsx-sample'},files={'file':(name,stream.getvalue())})
        assert r.status_code==201,r.text
        assert r.json()['data']['row_count']==1


def test_xlsx_missing_column(client):
    wb=Workbook();ws=wb.active;ws.append(['loan_id']);ws.append(['L001']);buf=io.BytesIO();wb.save(buf)
    r=client.post('/api/import/loans',data={'dataset_id':'bad'},files={'file':('bad.xlsx',buf.getvalue())})
    assert r.status_code==422


def test_reject_unsupported_extension(client):
    r=_policy(client,'bad.exe',b'ignore')
    assert r.status_code==422
